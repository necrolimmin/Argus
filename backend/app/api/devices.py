"""Devices tab: cameras and recorders, their channels, bulk add and import."""

import asyncio
import base64
import csv
import io
import uuid

from fastapi import APIRouter, HTTPException, Response, UploadFile, status
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import SessionLocal
from app.deps import AdminUser, Session
from app.models import Camera, CameraVendor, CredentialProfile, Device, DeviceKind, Direction, OrgUnit
from app.schemas import (
    URL_VENDORS,
    BulkCreate,
    BulkResult,
    ChannelIn,
    ChannelsIn,
    Connection,
    DeviceCreate,
    DeviceOut,
    DeviceUpdate,
    DuplicateOut,
    ImportIn,
    ImportRow,
    ProbeChannelOut,
    ProbeOut,
)
from app.security import decrypt_secret, encrypt_secret
from app.services.org import get_unit_in_scope
from app.services.probe import ProbeResult, ProbeTarget, probe
from app.services.streams import credentials, safe_register, safe_unregister

router = APIRouter(prefix="/devices", tags=["devices"])


# --- helpers -------------------------------------------------------------------


def device_out(dev: Device) -> DeviceOut:
    out = DeviceOut.model_validate(dev)
    out.has_password = bool(dev.password_enc) or dev.credential_id is not None
    return out


async def _load(session: AsyncSession, device_id: uuid.UUID) -> Device:
    dev = await session.get(Device, device_id, populate_existing=True)
    if dev is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "device_not_found")
    return dev


async def _resolve_credentials(session, credential_id, username: str, password: str) -> tuple[str, str]:
    if credential_id is None:
        return username, password
    profile = await session.get(CredentialProfile, credential_id)
    if profile is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "credential_not_found")
    return profile.username, decrypt_secret(profile.password_enc)


def _target(conn: Connection | Device, username: str, password: str, **extra) -> ProbeTarget:
    return ProbeTarget(
        vendor=conn.vendor,
        host=conn.host,
        rtsp_port=conn.rtsp_port,
        http_port=conn.http_port,
        username=username,
        password=password,
        **extra,
    )


async def _probe_out(session: AsyncSession, res: ProbeResult, device: Device | None = None) -> ProbeOut:
    """Adds what Argus already knows: duplicates and channels already added."""
    duplicates: list[DuplicateOut] = []
    conditions = []
    if res.serial:
        conditions.append(Device.serial == res.serial)
    if res.mac:
        conditions.append(Device.mac == res.mac)
    if conditions:
        query = select(Device).where(or_(*conditions))
        if device is not None:
            query = query.where(Device.id != device.id)
        for other in (await session.scalars(query)).unique():
            reason = "serial" if res.serial and other.serial == res.serial else "mac"
            duplicates.append(DuplicateOut(device_id=other.id, name=other.name, reason=reason))

    channel_ips = [c.ip for c in res.channels if c.ip]
    by_host = {}
    if channel_ips:
        rows = await session.scalars(select(Device).where(Device.host.in_(channel_ips)))
        by_host = {d.host: d.name for d in rows.unique()}
    existing = {c.channel: c.id for c in device.cameras} if device else {}

    return ProbeOut(
        ok=res.ok,
        error=res.error,
        kind=res.kind,
        name=res.name,
        model=res.model,
        serial=res.serial,
        mac=res.mac,
        firmware=res.firmware,
        codec=res.codec,
        rtsp_port=res.rtsp_port,
        channels=[
            ProbeChannelOut(
                channel=c.channel,
                name=c.name,
                rtsp_main=c.rtsp_main,
                rtsp_sub=c.rtsp_sub,
                ip=c.ip,
                camera_id=existing.get(c.channel),
                duplicate_of=by_host.get(c.ip, ""),
            )
            for c in res.channels
        ],
        snapshot="data:image/jpeg;base64," + base64.b64encode(res.snapshot).decode() if res.snapshot else None,
        warnings=res.warnings,
        duplicates=duplicates,
    )


def _camera(dev: Device, ch: ChannelIn) -> Camera:
    # Attached by the caller through dev.cameras.append().
    return Camera(
        name=ch.name.strip(),
        channel=ch.channel,
        rtsp_main=ch.rtsp_main,
        rtsp_sub=ch.rtsp_sub,
        org_unit_id=ch.org_unit_id or dev.org_unit_id,
    )


async def _register_all(dev: Device) -> None:
    for cam in dev.cameras:
        await safe_register(cam)


def _set_credentials(dev: Device, credential_id, username: str, password: str, keep_password: bool) -> None:
    dev.credential_id = credential_id
    if credential_id is not None:
        dev.username, dev.password_enc = "", ""
        return
    dev.username = username
    if password or not keep_password:
        dev.password_enc = encrypt_secret(password)


# --- list, check, create -------------------------------------------------------


@router.get("", response_model=list[DeviceOut])
async def list_devices(_: AdminUser, session: Session):
    devices = (await session.scalars(select(Device).order_by(Device.name))).unique().all()
    return [device_out(d) for d in devices]


@router.post("/probe", response_model=ProbeOut)
async def probe_device(body: Connection, _: AdminUser, session: Session):
    username, password = await _resolve_credentials(session, body.credential_id, body.username, body.password)
    res = await probe(_target(body, username, password, rtsp_main=body.rtsp_main, rtsp_sub=body.rtsp_sub,
                              channel_count=body.channel_count))
    return await _probe_out(session, res)


@router.post("", response_model=DeviceOut, status_code=status.HTTP_201_CREATED)
async def create_device(body: DeviceCreate, user: AdminUser, session: Session):
    await get_unit_in_scope(session, user, body.org_unit_id)
    for ch in body.channels:
        if ch.org_unit_id:
            await get_unit_in_scope(session, user, ch.org_unit_id)
    if body.credential_id:
        await _resolve_credentials(session, body.credential_id, "", "")
    if len({c.channel for c in body.channels}) != len(body.channels):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "duplicate_channel")

    dev = Device(
        name=body.name.strip(),
        kind=DeviceKind.recorder if len(body.channels) > 1 else body.kind,
        vendor=CameraVendor.rtsp if body.vendor == CameraVendor.other else body.vendor,
        host=body.host,
        rtsp_port=body.rtsp_port,
        http_port=body.http_port,
        org_unit_id=body.org_unit_id,
        model=body.model,
        serial=body.serial,
        mac=body.mac.lower(),
        firmware=body.firmware,
    )
    _set_credentials(dev, body.credential_id, body.username, body.password, keep_password=False)
    if body.vendor in URL_VENDORS and not dev.host:
        from urllib.parse import urlsplit

        dev.host = urlsplit(body.rtsp_main).hostname or ""
    for ch in body.channels:
        if body.vendor in URL_VENDORS and not ch.rtsp_main:
            ch.rtsp_main, ch.rtsp_sub = body.rtsp_main, body.rtsp_sub
        dev.cameras.append(_camera(dev, ch))
    session.add(dev)
    await session.commit()
    dev = await _load(session, dev.id)
    await _register_all(dev)
    return device_out(dev)


# --- bulk add (from a scan) ----------------------------------------------------


@router.post("/bulk", response_model=list[BulkResult])
async def bulk_create(body: BulkCreate, user: AdminUser, session: Session):
    """Checks every address and adds each device with all of its channels."""
    await get_unit_in_scope(session, user, body.org_unit_id)
    username, password = await _resolve_credentials(session, body.credential_id, body.username, body.password)
    limit = asyncio.Semaphore(8)

    async def one(t) -> tuple:
        async with limit:
            return t, await probe(_target(t, username, password))

    results: list[BulkResult] = []
    for t, res in await asyncio.gather(*(one(t) for t in body.targets)):
        if not res.ok:
            results.append(BulkResult(host=t.host, ok=False, error=res.error))
            continue
        async with SessionLocal() as s:
            dev = Device(
                name=(t.name or res.name or f"{res.model or t.vendor.value} {t.host}").strip()[:200],
                kind=res.kind,
                vendor=t.vendor,
                host=t.host,
                rtsp_port=res.rtsp_port or t.rtsp_port,
                http_port=t.http_port,
                org_unit_id=body.org_unit_id,
                model=res.model,
                serial=res.serial,
                mac=res.mac,
                firmware=res.firmware,
            )
            _set_credentials(dev, body.credential_id, body.username, body.password, keep_password=False)
            for c in res.channels:
                default = dev.name if len(res.channels) == 1 else f"{dev.name} — {c.channel}"
                dev.cameras.append(
                    _camera(dev, ChannelIn(channel=c.channel, name=(c.name or default)[:200],
                                           rtsp_main=c.rtsp_main, rtsp_sub=c.rtsp_sub))
                )
            s.add(dev)
            await s.commit()
            dev = await s.get(Device, dev.id, populate_existing=True)
            await _register_all(dev)
            results.append(BulkResult(host=t.host, ok=True, device_id=dev.id, cameras=len(dev.cameras)))
    return results


# --- import from Excel / CSV ---------------------------------------------------

TEMPLATE_HEADER = [
    "Название", "Производитель", "IP", "Порт RTSP", "Логин", "Пароль", "Канал",
    "Подразделение (Компания / Филиал / Предприятие / Раздел)", "Распознавание (да/нет)",
    "Направление (вход/выход)",
]
VENDOR_ALIASES = {v.value: v for v in CameraVendor} | {
    "hik": CameraVendor.hikvision, "tp-link": CameraVendor.vigi, "tplink": CameraVendor.vigi,
    "xmeye": CameraVendor.dvrip, "unv": CameraVendor.uniview,
}
YES = {"да", "yes", "y", "1", "true", "ha", "ҳа", "+"}
DIRECTIONS = {
    "вход": Direction.entry, "entry": Direction.entry, "in": Direction.entry, "kirish": Direction.entry,
    "кириш": Direction.entry, "выход": Direction.exit, "exit": Direction.exit, "out": Direction.exit,
    "chiqish": Direction.exit, "чиқиш": Direction.exit,
}


@router.get("/import/template")
async def import_template(_: AdminUser):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Cameras"
    ws.append(TEMPLATE_HEADER)
    ws.append(["Проходная — вход", "hikvision", "192.168.1.64", 554, "admin", "", 1,
               "Компания / Филиал Ташкент / Завод", "да", "вход"])
    ws.append(["Склад, ворота 1", "dahua", "192.168.1.108", 554, "admin", "", 1,
               "Компания / Филиал Ташкент / Завод / Склад", "нет", ""])
    for col, width in zip("ABCDEFGHIJ", (28, 16, 16, 10, 12, 12, 8, 48, 18, 18)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="argus-cameras.xlsx"'},
    )


def _read_rows(filename: str, data: bytes) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [["" if v is None else str(v).strip() for v in row] for row in wb.active.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


async def _unit_resolver(session):
    units = (await session.scalars(select(OrgUnit))).all()
    index = {(u.parent_id, u.name.strip().lower()): u.id for u in units}

    def resolve(path: str) -> uuid.UUID | None:
        parent = None
        for part in [p.strip().lower() for p in path.split("/") if p.strip()]:
            parent = index.get((parent, part))
            if parent is None:
                return None
        return parent

    return resolve


def _validate(row: ImportRow) -> ImportRow:
    errors = []
    if not row.name:
        errors.append("name_required")
    if row.vendor not in VENDOR_ALIASES:
        errors.append("vendor_unknown")
    elif VENDOR_ALIASES[row.vendor] in URL_VENDORS:
        if not row.host.lower().startswith(("rtsp://", "http://", "https://")):
            errors.append("url_required")
    else:
        import ipaddress

        try:
            ipaddress.ip_address(row.host)
        except ValueError:
            if not row.host or " " in row.host:
                errors.append("host_invalid")
    if row.org_unit_id is None:
        errors.append("unit_not_found")
    if not 1 <= row.channel <= 256:
        errors.append("channel_invalid")
    row.errors = errors
    return row


@router.post("/import/preview", response_model=list[ImportRow])
async def import_preview(file: UploadFile, _: AdminUser, session: Session):
    data = await file.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "file_too_large")
    try:
        table = _read_rows(file.filename or "", data)
    except Exception as exc:  # noqa: BLE001  (corrupt workbook, unknown format)
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "import_unreadable") from exc
    resolve = await _unit_resolver(session)
    rows = []
    for i, cells in enumerate(table[1:], start=2):
        cells = (cells + [""] * 10)[:10]
        if not any(cells):
            continue
        name, vendor, host, port, user, password, channel, org_path, face, direction = cells
        row = ImportRow(
            row=i,
            name=name,
            vendor=vendor.strip().lower(),
            host=host,
            port=int(float(port)) if port.replace(".", "", 1).isdigit() else None,
            username=user,
            password=password,
            channel=int(float(channel)) if channel.replace(".", "", 1).isdigit() else 1,
            org_path=org_path,
            org_unit_id=resolve(org_path),
            face_recognition=face.strip().lower() in YES,
        )
        row.direction = DIRECTIONS.get(direction.strip().lower(), Direction.none) if row.face_recognition else Direction.none
        rows.append(_validate(row))
    if len(rows) > 2000:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "import_too_many_rows")
    return rows


@router.post("/import")
async def import_devices(body: ImportIn, user: AdminUser, session: Session):
    """Rows with the same address and login become one device (a recorder's channels)."""
    resolve = await _unit_resolver(session)
    groups: dict[tuple, list[ImportRow]] = {}
    skipped = []
    for row in body.rows:
        row.org_unit_id = resolve(row.org_path)
        if _validate(row).errors:
            skipped.append(row.row)
            continue
        vendor = VENDOR_ALIASES[row.vendor]
        key = (vendor, row.host.strip(), row.port, row.username)
        groups.setdefault(key, []).append(row)

    created: list[Device] = []
    for (vendor, host, port, username), rows in groups.items():
        url = vendor in URL_VENDORS
        first = rows[0]
        dev = Device(
            name=first.name if len(rows) == 1 else f"{vendor.value} {host}",
            kind=DeviceKind.recorder if len(rows) > 1 else DeviceKind.camera,
            vendor=CameraVendor.rtsp if vendor == CameraVendor.other else vendor,
            host="" if url else host,
            rtsp_port=port or (34567 if vendor == CameraVendor.dvrip else 554),
            org_unit_id=first.org_unit_id,
        )
        if url:
            from urllib.parse import urlsplit

            dev.host = urlsplit(host).hostname or ""
        _set_credentials(dev, None, username, first.password, keep_password=False)
        seen = set()
        for r in rows:
            if r.channel in seen:
                skipped.append(r.row)
                continue
            seen.add(r.channel)
            cam = _camera(dev, ChannelIn(channel=r.channel, name=r.name, org_unit_id=r.org_unit_id,
                                         rtsp_main=host if url else ""))
            cam.face_recognition = r.face_recognition
            cam.direction = r.direction
            dev.cameras.append(cam)
        session.add(dev)
        created.append(dev)
    await session.commit()

    cameras = 0
    for dev in created:
        dev = await _load(session, dev.id)
        cameras += len(dev.cameras)
        await _register_all(dev)
    return {"devices": len(created), "cameras": cameras, "skipped_rows": sorted(skipped)}


# --- one device ----------------------------------------------------------------


@router.put("/{device_id}", response_model=DeviceOut)
async def update_device(device_id: uuid.UUID, body: DeviceUpdate, user: AdminUser, session: Session):
    dev = await _load(session, device_id)
    await get_unit_in_scope(session, user, body.org_unit_id)
    if body.credential_id:
        await _resolve_credentials(session, body.credential_id, "", "")
    dev.name = body.name.strip()
    dev.org_unit_id = body.org_unit_id
    if dev.vendor not in URL_VENDORS:
        dev.host = body.host.strip()
    dev.rtsp_port, dev.http_port = body.rtsp_port, body.http_port
    _set_credentials(dev, body.credential_id, body.username, body.password, keep_password=True)
    await session.commit()
    dev = await _load(session, device_id)
    await _register_all(dev)
    return device_out(dev)


@router.delete("/{device_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_device(device_id: uuid.UUID, _: AdminUser, session: Session):
    dev = await _load(session, device_id)
    camera_ids = [c.id for c in dev.cameras]
    await session.delete(dev)
    await session.commit()
    for cid in camera_ids:
        await safe_unregister(cid)


@router.post("/{device_id}/sync", response_model=ProbeOut)
async def sync_device(device_id: uuid.UUID, _: AdminUser, session: Session):
    """Re-reads the device: identity, status and (for recorders) its channel list.

    Channels that disappeared are flagged on their cameras; new ones are
    returned without `camera_id`, to be added with POST .../cameras.
    """
    dev = await _load(session, device_id)
    username, password = credentials(dev)
    first = dev.cameras[0] if dev.cameras else None
    res = await probe(_target(dev, username, password,
                              rtsp_main=first.rtsp_main if first else "",
                              rtsp_sub=first.rtsp_sub if first else ""))
    if res.ok:
        dev.model = res.model or dev.model
        dev.serial = res.serial or dev.serial
        dev.mac = res.mac or dev.mac
        dev.firmware = res.firmware or dev.firmware
        port_changed = bool(res.rtsp_port) and res.rtsp_port != dev.rtsp_port
        if port_changed:
            dev.rtsp_port = res.rtsp_port
        if "api_unavailable" not in res.warnings and dev.vendor not in URL_VENDORS:
            listed = {c.channel for c in res.channels}
            for cam in dev.cameras:
                cam.missing = cam.channel not in listed
        await session.commit()
        dev = await _load(session, device_id)
        if port_changed:
            await _register_all(dev)
    return await _probe_out(session, res, dev)


@router.post("/{device_id}/cameras", response_model=DeviceOut)
async def add_channels(device_id: uuid.UUID, body: ChannelsIn, user: AdminUser, session: Session):
    dev = await _load(session, device_id)
    taken = {c.channel for c in dev.cameras}
    for ch in body.channels:
        if ch.channel in taken:
            continue
        if ch.org_unit_id:
            await get_unit_in_scope(session, user, ch.org_unit_id)
        dev.cameras.append(_camera(dev, ch))
        taken.add(ch.channel)
    if len(dev.cameras) > 1:
        dev.kind = DeviceKind.recorder
    await session.commit()
    dev = await _load(session, device_id)
    await _register_all(dev)
    return device_out(dev)
